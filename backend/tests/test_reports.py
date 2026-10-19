"""Teacher review, approval and the exported report, against the local database."""

from decimal import Decimal
from io import BytesIO

import pymupdf
import pytest
from openpyxl import load_workbook

from app.core.db import transaction
from app.modules.scoring import score_job
from tests.conftest import OTHER_TEACHER, TEACHER_ID
from tests.fakes import FakeLanguageModel

SURE = {"verdicts": [{"key_point": 1, "verdict": "covered", "evidence": "stack"}], "confidence": 0.9, "reason": "Right."}
UNSURE = {**SURE, "confidence": 0.3}


@pytest.fixture
def finished(login, students, exam, monkeypatch):
    """A closed exam. Student A answered all four questions: Q1 scored, Q2 a wrong exact output,
    Q3 flagged (the AI was unsure), Q4 not scored yet. Student B answered Q1, then the exam closed."""
    model = FakeLanguageModel(reply=SURE)
    monkeypatch.setattr(score_job, "get_language_model", lambda: model)

    def sit(student, roll, answers):
        client = login(student)
        sid = client.post("/sessions", json={"join_code": exam["join_code"], "roll_number": roll,
                                             "experiment_id": exam["experiment_id"]}).json()["id"]
        client.post(f"/sessions/{sid}/start")
        ids = []
        for position, text in enumerate(answers, start=1):
            client.post(f"/sessions/{sid}/answers",
                        json={"position": position, "text": text, "idempotency_key": f"k-{sid}-{position}"})
            with transaction() as conn:
                ids.append(str(conn.execute(
                    """select a.id from public.answers a join public.session_questions sq on sq.id = a.session_question_id
                       where sq.session_id = %s and sq.position = %s""", (sid, position)).fetchone()["id"]))
        return ids

    a = sit(students[0], "22CS001", ["A stack is LIFO.", "5, 3", "stack =HYPERLINK(\"http://evil\")", "stack underflow"])
    b = sit(students[1], "22CS002", ["the stack top"])
    score_job.run({"answer_id": a[0]})
    score_job.run({"answer_id": a[1]})
    model.reply = UNSURE
    score_job.run({"answer_id": a[2]})
    model.reply = SURE
    score_job.run({"answer_id": b[0]})
    login(TEACHER_ID).post(f"/exams/{exam['id']}/close")
    return {"exam": exam["id"], "a": a, "b": b, "teacher": login(TEACHER_ID)}


def answers_of(sheet, roll):
    return next(s for s in sheet["students"] if s["roll_number"] == roll)["answers"]


def test_the_sheet_shows_each_answers_mark_and_what_still_needs_the_teacher(finished):
    sheet = finished["teacher"].get(f"/exams/{finished['exam']}/sheet").json()

    a = answers_of(sheet, "22CS001")
    assert [x["status"] for x in a] == ["scored", "scored", "to_check", "scoring"]
    assert [x["final"] for x in a] == [4.0, 0.0, 4.0, None]
    assert a[2]["review_reason"].startswith("The AI was unsure")
    b = answers_of(sheet, "22CS002")
    assert [x["status"] for x in b] == ["scored", "unanswered", "unanswered", "unanswered"]
    assert next(s for s in sheet["students"] if s["roll_number"] == "22CS002")["total"] == 4.0
    assert (sheet["max_marks"], sheet["waiting"]) == (16.0, 2)


def test_a_teachers_mark_sits_beside_the_ais_and_wins(finished):
    teacher, answer = finished["teacher"], finished["a"][2]

    response = teacher.put(f"/answers/{answer}/review", json={"teacher_score": 2.5, "note": "Only half right."})

    assert response.status_code == 200
    with transaction() as conn:
        score = conn.execute("select * from public.scores where answer_id = %s", (answer,)).fetchone()
    assert (score["ai_score"], score["teacher_score"]) == (Decimal("4.00"), Decimal("2.50"))  # both kept
    assert score["reason"] == "Right." and score["teacher_note"] == "Only half right."
    assert str(score["reviewed_by"]) == TEACHER_ID
    row = answers_of(teacher.get(f"/exams/{finished['exam']}/sheet").json(), "22CS001")[2]
    assert (row["status"], row["final"]) == ("reviewed", 2.5)


def test_accepting_the_ais_mark_marks_it_reviewed(finished):
    teacher = finished["teacher"]

    teacher.put(f"/answers/{finished['a'][2]}/review", json={"teacher_score": None})

    row = answers_of(teacher.get(f"/exams/{finished['exam']}/sheet").json(), "22CS001")[2]
    assert (row["status"], row["final"], row["teacher_score"]) == ("reviewed", 4.0, None)


def test_an_unscored_answer_can_be_marked_by_the_teacher(finished):
    teacher = finished["teacher"]
    # There is no AI mark to accept yet, so the teacher must give one.
    assert teacher.put(f"/answers/{finished['a'][3]}/review", json={"teacher_score": None}).status_code == 422

    teacher.put(f"/answers/{finished['a'][3]}/review", json={"teacher_score": 3})

    row = answers_of(teacher.get(f"/exams/{finished['exam']}/sheet").json(), "22CS001")[3]
    assert (row["status"], row["final"], row["ai_score"]) == ("reviewed", 3.0, None)


def test_marks_are_checked_and_belong_to_the_exams_teacher(finished, login):
    teacher, answer = finished["teacher"], finished["a"][0]

    assert teacher.put(f"/answers/{answer}/review", json={"teacher_score": 4.5}).status_code == 422
    assert teacher.put(f"/answers/{answer}/review", json={"teacher_score": -1}).status_code == 422
    assert login(OTHER_TEACHER).put(f"/answers/{answer}/review", json={"teacher_score": 1}).status_code == 404
    assert login(OTHER_TEACHER).get(f"/exams/{finished['exam']}/sheet").status_code == 404


def test_marks_cannot_be_reviewed_while_the_exam_is_open(login, students, exam):
    client = login(students[0])
    sid = client.post("/sessions", json={"join_code": exam["join_code"], "roll_number": "22CS001",
                                         "experiment_id": exam["experiment_id"]}).json()["id"]
    client.post(f"/sessions/{sid}/start")
    client.post(f"/sessions/{sid}/answers", json={"position": 1, "text": "x", "idempotency_key": "k-open-1"})
    with transaction() as conn:
        answer = conn.execute(
            "select a.id from public.answers a where a.idempotency_key = 'k-open-1'").fetchone()["id"]

    assert login(TEACHER_ID).put(f"/answers/{answer}/review", json={"teacher_score": 1}).status_code == 409
    assert login(TEACHER_ID).post(f"/exams/{exam['id']}/approve").status_code == 409


def test_approval_waits_for_every_answer_to_have_a_mark_then_locks_them(finished):
    teacher, exam = finished["teacher"], finished["exam"]

    refused = teacher.post(f"/exams/{exam}/approve")
    assert refused.status_code == 409
    assert "2 answer(s) still need a mark: 22CS001 Q3, 22CS001 Q4" in refused.json()["detail"]

    teacher.put(f"/answers/{finished['a'][2]}/review", json={"teacher_score": None})
    teacher.put(f"/answers/{finished['a'][3]}/review", json={"teacher_score": 2})
    assert teacher.post(f"/exams/{exam}/approve").json()["state"] == "approved"

    assert teacher.put(f"/answers/{finished['a'][0]}/review", json={"teacher_score": 1}).status_code == 409
    assert teacher.post(f"/exams/{exam}/approve").status_code == 200  # approving twice is harmless


@pytest.fixture
def approved(finished):
    teacher = finished["teacher"]
    teacher.put(f"/answers/{finished['a'][2]}/review", json={"teacher_score": 2.5, "note": "Half."})
    teacher.put(f"/answers/{finished['a'][3]}/review", json={"teacher_score": 2})
    assert teacher.post(f"/exams/{finished['exam']}/approve").status_code == 200
    return finished


def test_the_report_is_exported_only_once_approved(finished):
    assert finished["teacher"].get(f"/exams/{finished['exam']}/report.xlsx").status_code == 409
    assert finished["teacher"].get(f"/exams/{finished['exam']}/report.pdf").status_code == 409


def test_the_excel_report_has_both_marks_and_never_runs_student_text(approved):
    response = approved["teacher"].get(f"/exams/{approved['exam']}/report.xlsx")

    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="Test-lab-exam-A.xlsx"'
    book = load_workbook(BytesIO(response.content))
    marks = [[c.value for c in row] for row in book["Class sheet"].iter_rows(values_only=False)]
    assert marks[0] == ["Roll number", "Experiment", "Q1", "Q2", "Q3", "Q4", "Total", "Out of"]
    assert marks[1] == ["22CS001", "Stack Using Arrays", 4, 0, 2.5, 2, 8.5, 16]
    answers = list(book["Answers"].iter_rows(min_row=2))
    q3 = answers[2]
    assert (q3[7].value, q3[11].value, q3[12].value, q3[13].value) == (4, 2.5, "Half.", 2.5)  # AI, teacher, note, final
    assert q3[4].value == 'stack =HYPERLINK("http://evil")' and q3[4].data_type == "s"


def test_the_pdf_report_has_the_class_sheet_and_each_students_answers(approved):
    response = approved["teacher"].get(f"/exams/{approved['exam']}/report.pdf")

    assert response.status_code == 200
    text = "".join(page.get_text() for page in pymupdf.open(stream=response.content, filetype="pdf"))
    assert "22CS001" in text and "22CS002" in text
    assert "8.5 / 16" in text
    assert "Teacher: 2.5 - Half." in text
    assert "Answer: (not reached)" in text
