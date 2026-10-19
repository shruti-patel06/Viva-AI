"""The approved report as Excel and PDF, both built from the same Sheet.

Student answers are untrusted input here too: an answer starting with "=" would become a live
formula in Excel, so every text cell is written as plain text.
"""

from datetime import datetime
from decimal import Decimal
from io import BytesIO

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font

from app.modules.reports.sheet import AnswerRow, Sheet

MARKING_RULE = (
    "Final mark = the teacher's mark where one was set, otherwise the AI's suggested mark. "
    "AI marks: each key point covered counts 1, partly 0.5, missed 0, scaled to the question's marks, "
    "less the hint penalty for each hint taken. An unanswered question scores 0."
)


def _num(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _verdicts(answer: AnswerRow) -> str:
    return "\n".join(
        f"{hit['verdict']}: {hit['key_point']}" + (f' ("{hit["evidence"]}")' if hit.get("evidence") else "")
        for hit in answer.key_points_hit or []
    )


def _mark(value: Decimal | None) -> str:
    """8.50 -> "8.5", 4.00 -> "4"; no mark -> "-"."""
    return "-" if value is None else f"{float(value):g}"


def _when(value: datetime | None) -> str:
    return value.strftime("%Y-%m-%d %H:%M") if value else ""


# ---- Excel ----------------------------------------------------------------------------


def _row(sheet, values: list) -> None:
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        _as_text(cell)


def _as_text(cell: Cell) -> None:
    # openpyxl turns any string starting with "=" into a formula. Student text must never run.
    if isinstance(cell.value, str):
        cell.data_type = "s"


def _header(sheet, titles: list[str]) -> None:
    sheet.append(titles)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    # By name, not sheet.cell(): asking for the cell would create an empty row under the header.
    sheet.freeze_panes = f"A{sheet.max_row + 1}"


def to_xlsx(report: Sheet) -> bytes:
    book = Workbook()
    count = max((len(s.answers) for s in report.students), default=0)

    marks = book.active
    marks.title = "Class sheet"
    _header(marks, ["Roll number", "Experiment", *[f"Q{n}" for n in range(1, count + 1)], "Total", "Out of"])
    for student in report.students:
        _row(marks, [student.roll_number, student.experiment, *[_num(a.final) for a in student.answers],
                     _num(student.total), _num(report.max_marks)])
    marks.column_dimensions["B"].width = 28

    answers = book.create_sheet("Answers")
    _header(answers, [
        "Roll number", "Q", "Type", "Question", "Answer", "Speech to text heard", "Hints", "AI score",
        "AI reason", "Key points", "Flagged because", "Teacher score", "Teacher note", "Final", "Reviewed by",
        "Reviewed at", "Prompt version", "Model",
    ])
    for student in report.students:
        for a in student.answers:
            _row(answers, [
                student.roll_number, a.position, a.type + (" (generated live)" if a.live else ""), a.question, (a.text or "(skipped)") if a.answer_id else "(not reached)",
                a.heard or "", a.hints_used, _num(a.ai_score), a.ai_reason or "", _verdicts(a), a.review_reason or "",
                _num(a.teacher_score), a.teacher_note or "", _num(a.final), a.reviewed_by or "", _when(a.reviewed_at),
                a.prompt_version or "", a.model or "",
            ])
    for column, width in {"D": 40, "E": 40, "F": 30, "I": 40, "J": 50, "K": 30, "M": 30}.items():
        answers.column_dimensions[column].width = width
    for row in answers.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    about = book.create_sheet("About")
    for label, value in [
        ("Exam", report.title), ("Subject", report.subject), ("Section", report.section),
        ("Approved by", report.approved_by or ""), ("Approved at", _when(report.approved_at)),
        ("Marks per question", _num(report.marks_per_question)),
        ("Hint penalty", f"{float(report.hint_penalty) * 100:g}% per hint"), ("Marking rule", MARKING_RULE),
    ]:
        _row(about, [label, value])
    about.column_dimensions["A"].width = 20
    about.column_dimensions["B"].width = 90

    out = BytesIO()
    book.save(out)
    return out.getvalue()


# ---- PDF ------------------------------------------------------------------------------


def _latin(text: str | None) -> str:
    """The PDF's built-in fonts cover Latin-1 only. Common punctuation is swapped for plain
    equivalents; anything else becomes "?" rather than failing the export."""
    if not text:
        return ""
    for fancy, plain in {"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-",
                         "…": "...", "₹": "Rs ", "→": "->"}.items():
        text = text.replace(fancy, plain)
    return text.encode("latin-1", "replace").decode("latin-1")


class _Report(FPDF):
    def __init__(self, title: str):
        super().__init__()
        self.report_title = title
        self.set_auto_page_break(auto=True, margin=15)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Helvetica", size=8)
        self.cell(0, 8, _latin(f"{self.report_title}  -  page {self.page_no()}"), align="C")


def _line(pdf: FPDF, text: str, size: int = 10, bold: bool = False, gap: float = 1) -> None:
    pdf.set_font("Helvetica", style="B" if bold else "", size=size)
    pdf.multi_cell(0, size * 0.5, _latin(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(gap)


def to_pdf(report: Sheet) -> bytes:
    pdf = _Report(f"{report.title} ({report.section})")
    pdf.add_page()
    _line(pdf, report.title, size=16, bold=True)
    _line(pdf, f"{report.subject}, section {report.section}")
    _line(pdf, f"Approved by {report.approved_by or '-'} on {_when(report.approved_at) or '-'}", gap=3)
    _line(pdf, MARKING_RULE, size=8, gap=4)

    count = max((len(s.answers) for s in report.students), default=0)
    pdf.set_font("Helvetica", size=9)
    with pdf.table(col_widths=(30, 50, *([12] * count), 20), text_align="LEFT") as table:
        table.row(["Roll number", "Experiment", *[f"Q{n}" for n in range(1, count + 1)], "Total"])
        for s in report.students:
            table.row([_latin(s.roll_number), _latin(s.experiment), *[_mark(a.final) for a in s.answers],
                       f"{_mark(s.total)} / {_mark(report.max_marks)}"])

    for student in report.students:
        pdf.add_page()
        _line(pdf, f"{student.roll_number} - {student.experiment}", size=13, bold=True)
        _line(pdf, f"Total: {_mark(student.total)} out of {_mark(report.max_marks)}", gap=4)
        for a in student.answers:
            label = a.type.replace("_", " ") + (", generated live" if a.live else "")
            _line(pdf, f"Q{a.position} ({label}): {a.question}", bold=True)
            given = "(not reached)" if a.answer_id is None else (a.text or "(skipped)")
            _line(pdf, f"Answer: {given}")
            if a.heard and a.heard != a.text:
                _line(pdf, f"Speech to text heard: {a.heard}", size=8)
            if a.hints_used:
                _line(pdf, f"Hints taken: {a.hints_used}", size=8)
            if a.answer_id:
                _line(pdf, f"AI: {_mark(a.ai_score)} - {a.ai_reason or a.review_reason or ''}", size=9)
                if a.teacher_score is not None:
                    _line(pdf, f"Teacher: {_mark(a.teacher_score)}" + (f" - {a.teacher_note}" if a.teacher_note else ""), size=9)
            _line(pdf, f"Final: {_mark(a.final)} / {_mark(report.marks_per_question)}", size=9, bold=True, gap=4)
    return bytes(pdf.output())
